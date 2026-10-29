"""Sara dokumentumkészítő eszközei.

- Szöveges dokumentum: Markdown a forrás, letöltéskor készül belőle Word (.docx) vagy PDF.
- Táblázat: strukturált JSON-ból Excel (.xlsx), képletekkel és számformátumokkal.
- Prezentáció: strukturált JSON-ból PowerPoint (.pptx).
- Olvasás: a tárolt .xlsx / .pptx / .docx szöveggé alakítása, hogy Sara lássa a tartalmukat.
- Szövegkinyerés a kereséshez: szöveges fájl, PDF és a fenti irodai formátumok.

Csak sima pip-csomagok kellenek: markdown-it-py, python-docx, reportlab, openpyxl, python-pptx, pypdf.
A PDF-hez a fonts/ mappában lévő DejaVu betűtípusok kellenek (az ő és ű betűk miatt).
"""
import datetime
import io
import os
import re
from html.parser import HTMLParser
from xml.sax.saxutils import escape

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
MD_MIME = "text/markdown"

MAX_SHEETS, MAX_ROWS, MAX_COLS = 10, 5000, 50
MAX_SLIDES, MAX_TABLE_ROWS, MAX_TABLE_COLS = 40, 15, 8
MAX_TEXT_FOR_CHAT = 150000  # ennyi karaktert adunk át Sarának egy-egy xlsx/pptx/docx tartalmából (kb. 60 oldal)


class ToolError(Exception):
    """Érthető hibaüzenet, amit Sara visszakap, és ki tud javítani."""


# ---------------------------------------------------------------------------
# Eszközdefiníciók Claude számára
# ---------------------------------------------------------------------------
TOOLS = [
    {
        "name": "create_document",
        "description": (
            "Szöveges dokumentum (szerződés, indoklás, vélemény, levél, összefoglaló, jegyzet stb.) készítése "
            "Markdown formátumban. A dokumentum a projekt dokumentumtárába kerül, ahonnan a felhasználó Word, PDF "
            "vagy Markdown formátumban töltheti le. Csak akkor használd, ha a felhasználó dokumentum vagy fájl "
            "készítését kéri; egy egyszerű kérdésre a chatben válaszolj."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A dokumentum címe (a fájlnév is ebből készül)."},
                "content_markdown": {
                    "type": "string",
                    "description": (
                        "A teljes dokumentum Markdownban: # címsorok, listák, **kiemelés**, táblázatok "
                        "(| oszlop | oszlop |). A címet ne ismételd meg a tartalom elején, azt az alkalmazás beteszi."
                    ),
                },
            },
            "required": ["title", "content_markdown"],
        },
    },
    {
        "name": "create_spreadsheet",
        "description": (
            "Excel táblázat (.xlsx) készítése egy vagy több munkalappal. Az 1. sor a fejléc, az adatok a 2. sortól "
            "indulnak, az A oszlop az 1. oszlop. Képletet úgy adj meg, hogy a cella értéke '='-lel kezdődő szöveg, "
            "angol függvénynevekkel és vesszős argumentum-elválasztással (pl. '=B2*C2', '=SUM(B2:B10)', "
            "'=IF(B2>0,B2*0.27,0)'); az Excel magyarul jeleníti meg. A számokat számként add meg (ne szövegként), "
            "a dátumokat 'ÉÉÉÉ-HH-NN' szövegként. Százalékot törtként (0.27) és '0.0%' formátummal. A sorok "
            "száma ne legyen több 5000-nél. A diagramot nem kell elkészíteni."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A fájl neve."},
                "sheets": {
                    "type": "array",
                    "minItems": 1,
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string", "description": "Munkalap neve (max. 31 karakter)."},
                            "columns": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "header": {"type": "string"},
                                        "format": {
                                            "type": "string",
                                            "description": (
                                                "Excel számformátum az oszlop celláira, pl. '#,##0' (egész, "
                                                "ezres tagolással), '#,##0.00', '0.0%', 'yyyy-mm-dd', "
                                                "'#,##0 \"Ft\"', '@' (szöveg). Elhagyható."
                                            ),
                                        },
                                        "width": {"type": "number", "description": "Oszlopszélesség (opcionális)."},
                                    },
                                    "required": ["header"],
                                },
                            },
                            "rows": {
                                "type": "array",
                                "items": {
                                    "type": "array",
                                    "items": {"type": ["string", "number", "boolean", "null"]},
                                },
                                "description": "Az adatsorok; minden sor az oszlopok sorrendjében tartalmazza az értékeket.",
                            },
                        },
                        "required": ["name", "columns", "rows"],
                    },
                },
            },
            "required": ["title", "sheets"],
        },
    },
    {
        "name": "create_presentation",
        "description": (
            "PowerPoint prezentáció (.pptx) készítése. Dianként egy gondolat: rövid cím, legfeljebb 6-7 rövid "
            "felsorolási pont, vagy egy legfeljebb 15 soros, 8 oszlopos táblázat. Az első dia lehet címdia "
            "(csak 'title' és 'subtitle'). Az előadói jegyzet a 'notes' mezőbe kerül."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A prezentáció címe (a fájlnév is ebből készül)."},
                "slides": {
                    "type": "array",
                    "minItems": 1,
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "subtitle": {"type": "string", "description": "Csak címdián."},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                            "table": {
                                "type": "object",
                                "properties": {
                                    "columns": {"type": "array", "items": {"type": "string"}},
                                    "rows": {
                                        "type": "array",
                                        "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
                                    },
                                },
                                "required": ["columns", "rows"],
                            },
                            "notes": {"type": "string", "description": "Előadói jegyzet."},
                        },
                        "required": ["title"],
                    },
                },
            },
            "required": ["title", "slides"],
        },
    },
]


def safe_filename(title, ext):
    import unicodedata
    title = unicodedata.normalize("NFC", str(title or ""))   # a szétbontott ékezetek ne vesszenek el
    stem = re.sub(r"[^\w\- ]", "", str(title or "dokumentum"), flags=re.UNICODE).strip()
    stem = re.sub(r"\s+", "_", stem)[:80] or "dokumentum"
    return f"{stem}.{ext}"


def run_tool(name, tool_input):
    """Végrehajt egy eszközhívást. Visszaad: (fájlnév, mime, bájtok). Hibánál ToolError."""
    if not isinstance(tool_input, dict):
        raise ToolError("Hibás eszközbemenet.")
    title = str(tool_input.get("title") or "").strip()
    if name == "create_document":
        body = str(tool_input.get("content_markdown") or "").strip()
        if not body:
            raise ToolError("A content_markdown üres.")
        if not re.match(r"^#\s", body) and title:
            body = f"# {title}\n\n{body}"
        return safe_filename(title, "md"), MD_MIME, (body + "\n").encode("utf-8")
    if name == "create_spreadsheet":
        return safe_filename(title, "xlsx"), XLSX_MIME, build_xlsx(tool_input)
    if name == "create_presentation":
        return safe_filename(title, "pptx"), PPTX_MIME, build_pptx(tool_input)
    raise ToolError(f"Ismeretlen eszköz: {name}")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
_FORMAT_ALIASES = {
    "text": "@", "szöveg": "@", "szoveg": "@",
    "integer": "#,##0", "egész": "#,##0", "egesz": "#,##0",
    "number": "#,##0.00", "szám": "#,##0.00", "szam": "#,##0.00",
    "percent": "0.0%", "százalék": "0.0%", "szazalek": "0.0%",
    "date": "yyyy-mm-dd", "dátum": "yyyy-mm-dd", "datum": "yyyy-mm-dd",
    "currency": '#,##0 "Ft"', "pénz": '#,##0 "Ft"', "penz": '#,##0 "Ft"',
}
_BAD_FORMULA = re.compile(
    r"(WEBSERVICE|HYPERLINK|IMPORTDATA|IMPORTXML|IMPORTHTML|IMPORTRANGE|RTD|DDE|EXEC|CALL|REGISTER|URL)\s*\(|"
    r"https?:|file:|ftp:|\[|\\\\",
    re.IGNORECASE,
)
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _norm_format(fmt):
    if not fmt:
        return None
    fmt = str(fmt).strip()
    return _FORMAT_ALIASES.get(fmt.lower(), fmt)


def _is_date_format(fmt):
    if not fmt:
        return False
    return bool(re.search(r"[ymd]", re.sub(r'"[^"]*"', "", fmt).lower()))


def _prepare_value(value, fmt):
    if isinstance(value, str):
        if value.startswith("="):
            if _BAD_FORMULA.search(value):
                raise ToolError(f"A képlet nem engedélyezett (külső hivatkozás vagy tiltott függvény): {value[:60]}")
            return value
        if _is_date_format(fmt) and _ISO_DATE.match(value):
            try:
                return datetime.date.fromisoformat(value)
            except ValueError:
                return value
    return value


def _sheet_name(raw, used):
    name = re.sub(r"[\[\]:*?/\\]", " ", str(raw or "")).strip().strip("'")[:31] or "Munkalap"
    base, n = name, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def build_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ToolError("Legalább egy munkalap (sheets) kell.")
    if len(sheets) > MAX_SHEETS:
        raise ToolError(f"Legfeljebb {MAX_SHEETS} munkalap lehet.")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    side = Side(style="thin", color="BFBFBF")
    border = Border(left=side, right=side, top=side, bottom=side)

    for sh in sheets:
        cols = sh.get("columns") or []
        rows = sh.get("rows") or []
        if not cols:
            raise ToolError(f"A(z) '{sh.get('name', '')}' munkalapnak nincs oszlopa.")
        if len(cols) > MAX_COLS:
            raise ToolError(f"Legfeljebb {MAX_COLS} oszlop lehet munkalaponként.")
        if len(rows) > MAX_ROWS:
            raise ToolError(f"Legfeljebb {MAX_ROWS} sor lehet munkalaponként.")

        ws = wb.create_sheet(_sheet_name(sh.get("name"), used))
        fmts = [_norm_format(c.get("format")) for c in cols]
        lens = [len(str(c.get("header", ""))) for c in cols]

        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=str(col.get("header", "")))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = border

        for ri, row in enumerate(rows, 2):
            if not isinstance(row, list):
                raise ToolError(f"A(z) {ri - 1}. adatsor nem lista.")
            if len(row) > len(cols):
                raise ToolError(
                    f"A(z) {ri - 1}. adatsorban több érték ({len(row)}) van, mint oszlop ({len(cols)})."
                )
            for ci in range(len(cols)):
                value = _prepare_value(row[ci] if ci < len(row) else None, fmts[ci])
                cell = ws.cell(row=ri, column=ci + 1, value=value)
                if fmts[ci]:
                    cell.number_format = fmts[ci]
                cell.border = border
                is_text = isinstance(value, str) and not value.startswith("=")
                cell.alignment = Alignment(vertical="top", wrap_text=is_text)
                lens[ci] = max(lens[ci], len(str(value)) if is_text or isinstance(value, (int, float)) else 12)

        for ci, col in enumerate(cols, 1):
            width = col.get("width")
            if not isinstance(width, (int, float)) or width <= 0:
                width = min(max(lens[ci - 1] + 2, 8), 60)
            ws.column_dimensions[get_column_letter(ci)].width = width

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    wb.calculation.fullCalcOnLoad = True  # a képletek megnyitáskor kiszámolódjanak
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def build_pptx(spec):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ToolError("Legalább egy dia (slides) kell.")
    if len(slides) > MAX_SLIDES:
        raise ToolError(f"Legfeljebb {MAX_SLIDES} dia lehet.")

    navy, grey = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x59, 0x59, 0x59)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = str(spec.get("title") or "")

    def set_text(shape, text, size, bold=False, color=None):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.text = str(text)
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold = Pt(size), bold
                if color is not None:
                    r.font.color.rgb = color

    def place(shape, left, top, width, height):
        shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)

    def col_weights(columns, rows):
        w = []
        for i, c in enumerate(columns):
            longest = max([len(str(c))] + [len(str(r[i])) for r in rows if i < len(r) and r[i] is not None])
            w.append(max(min(longest, 40), 6))
        return w

    for idx, s in enumerate(slides):
        title = str(s.get("title") or "")
        bullets = [str(b) for b in (s.get("bullets") or [])]
        table = s.get("table")
        subtitle = s.get("subtitle")

        if idx == 0 and subtitle and not bullets and not table:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            place(slide.shapes.title, 0.8, 2.3, 11.7, 1.6)
            set_text(slide.shapes.title, title, 40, True, navy)
            place(slide.placeholders[1], 0.8, 4.1, 11.7, 1.2)
            set_text(slide.placeholders[1], subtitle, 22, False, grey)
            for p in slide.shapes.title.text_frame.paragraphs + slide.placeholders[1].text_frame.paragraphs:
                p.alignment = 1  # balra
        else:
            use_table = isinstance(table, dict)
            slide = prs.slides.add_slide(prs.slide_layouts[5 if use_table else 1])
            place(slide.shapes.title, 0.8, 0.4, 11.7, 0.95)
            set_text(slide.shapes.title, title, 30, True, navy)
            for p in slide.shapes.title.text_frame.paragraphs:
                p.alignment = 1
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(11.7), Inches(0.05))
            bar.fill.solid()
            bar.fill.fore_color.rgb = navy
            bar.line.fill.background()

            top = 1.7
            if use_table:
                cols = [str(c) for c in (table.get("columns") or [])]
                rows = table.get("rows") or []
                if not cols:
                    raise ToolError(f"A(z) {idx + 1}. dia táblázatának nincs oszlopa.")
                if len(cols) > MAX_TABLE_COLS or len(rows) > MAX_TABLE_ROWS:
                    raise ToolError(
                        f"A(z) {idx + 1}. dia táblázata túl nagy (legfeljebb {MAX_TABLE_ROWS} sor és "
                        f"{MAX_TABLE_COLS} oszlop fér el egy dián); oszd több diára."
                    )
                if bullets:
                    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(0.55 * len(bullets)))
                    box.text_frame.word_wrap = True
                    for i, b in enumerate(bullets):
                        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                        p.text = f"• {b}"
                        for r in p.runs:
                            r.font.size = Pt(20)
                    top += 0.55 * len(bullets) + 0.2
                n_rows = len(rows) + 1
                font = 18 if n_rows <= 6 else 16 if n_rows <= 9 else 14 if n_rows <= 12 else 11
                height = min(0.5 * n_rows, 7.5 - top - 0.5)
                gs = slide.shapes.add_table(n_rows, len(cols), Inches(0.8), Inches(top), Inches(11.7), Inches(height))
                tbl = gs.table
                weights = col_weights(cols, rows)
                total = Inches(11.7)
                for i, w in enumerate(weights):
                    tbl.columns[i].width = int(total * w / sum(weights))
                for ci, c in enumerate(cols):
                    tbl.cell(0, ci).text = c
                for ri, r in enumerate(rows, 1):
                    for ci in range(len(cols)):
                        v = r[ci] if ci < len(r) and r[ci] is not None else ""
                        tbl.cell(ri, ci).text = str(v)
                for ri_, r_ in enumerate(tbl.rows):
                    for cell in r_.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(font)
                                if ri_ == 0:
                                    run.font.bold = True
            else:
                body = slide.placeholders[1]
                place(body, 0.8, top, 11.7, 5.2)
                size = 26 if len(bullets) <= 4 else 24 if len(bullets) <= 6 else 20
                body.text_frame.word_wrap = True
                for i, b in enumerate(bullets or [""]):
                    p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
                    p.text = b
                    for r in p.runs:
                        r.font.size = Pt(size)

        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markdown -> blokkok (közös köztes forma a Word és a PDF számára)
# ---------------------------------------------------------------------------
_YEAR_AT_START = re.compile(r"^(\s*(?:[-*+]\s+|>\s*)*)(\d{4})\.(?=\s)", re.M)


def _protect_dates(text):
    """A magyar dátum („2026. 10. 15.”) sor elején a Markdown szerint számozott lista lenne (2026., 10., 15.
    egymásba ágyazva), és a dátum eltűnne. Az évszám pontját ezért escape-eljük; kódblokkban nem."""
    out, fence = [], False
    for line in text.split("\n"):
        if line.lstrip().startswith(("```", "~~~")):
            fence = not fence
        out.append(line if fence else _YEAR_AT_START.sub(r"\1\2\\.", line))
    return "\n".join(out)


def _md_to_html(text):
    from markdown_it import MarkdownIt

    # breaks: a sima sortörés sortörés marad (aláírásblokk, cím), ahogy a .md-ben látszik
    return MarkdownIt("commonmark", {"html": False, "breaks": True}).enable("table").render(_protect_dates(text))


class _Blocks(HTMLParser):
    """HTML -> blokklista. Blokk: ("h", szint, runs) | ("p", runs) | ("quote", runs) |
    ("li", mélység, rendezett, sorszám, runs) | ("code", runs) | ("hr",) | ("table", sorok).
    Run: (szöveg, félkövér, dőlt, kód). Táblázatcella: (runs, fejléc, igazítás)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.kind, self.extra, self.runs = [], None, {}, None
        self.bold = self.ital = self.code = self.quote = 0
        self.lists, self.in_pre = [], False
        self.table = self.row = self.cell = None
        self.cell_head, self.cell_align, self.href = False, None, None

    # -- segédek --
    def _start(self, kind, **extra):
        self._flush()
        self.kind, self.extra, self.runs = kind, extra, []

    def _flush(self):
        if self.kind is None:
            return
        runs = self.runs or []
        if runs and not self.in_pre:
            runs[0] = (runs[0][0].lstrip(),) + runs[0][1:]
            runs[-1] = (runs[-1][0].rstrip(),) + runs[-1][1:]
        if not self.in_pre:   # sortörés körül ne maradjon szóköz (a következő sor ne kezdődjön behúzva)
            for j in range(1, len(runs)):
                if runs[j - 1][0].endswith("\n"):
                    runs[j] = (runs[j][0].lstrip(" "),) + runs[j][1:]
        runs = [r for r in runs if r[0]]
        k = self.kind
        if k == "h" and runs:
            self.blocks.append(("h", self.extra["level"], runs))
        elif k in ("p", "quote") and runs:
            self.blocks.append((k, runs))
        elif k == "li" and runs:
            self.blocks.append(("li", self.extra["depth"], self.extra["ordered"], self.extra["n"], runs))
        elif k == "code":
            self.blocks.append(("code", [(("".join(r[0] for r in runs)).rstrip("\n"), False, False, True)]))
        self.kind, self.runs = None, None

    def _add(self, text, bold=None, ital=None, code=None):
        run = (text, bool(self.bold) if bold is None else bold, bool(self.ital) if ital is None else ital,
               bool(self.code) if code is None else code)
        if self.cell is not None:
            self.cell.append(run)
        else:
            if self.runs is None:
                self._start("quote" if self.quote else "p")
            self.runs.append(run)

    # -- HTML események --
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if re.fullmatch(r"h[1-6]", tag):
            self._start("h", level=int(tag[1]))
        elif tag == "p":
            if self.cell is None and self.kind != "li":
                self._start("quote" if self.quote else "p")
            elif self.kind == "li" and self.runs:
                self._add(" ")
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append([tag == "ol", int(a.get("start") or 1) - 1])
        elif tag == "li":
            entry = self.lists[-1] if self.lists else [False, 0]
            entry[1] += 1
            self._start("li", depth=max(len(self.lists), 1), ordered=entry[0], n=entry[1])
        elif tag in ("strong", "b"):
            self.bold += 1
        elif tag in ("em", "i"):
            self.ital += 1
        elif tag == "code" and not self.in_pre:
            self.code += 1
        elif tag == "pre":
            self._flush()
            self.in_pre, self.kind, self.runs = True, "code", []
        elif tag == "blockquote":
            self._flush()
            self.quote += 1
        elif tag == "br":
            self._add("\n")
        elif tag == "hr":
            self._flush()
            self.blocks.append(("hr",))
        elif tag == "a":
            self.href = a.get("href")
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr":
            self.row = []
        elif tag in ("th", "td"):
            self.cell, self.cell_head = [], tag == "th"
            m = re.search(r"text-align:\s*(\w+)", a.get("style") or "")
            self.cell_align = m.group(1) if m else None

    def handle_endtag(self, tag):
        if re.fullmatch(r"h[1-6]", tag):
            if self.kind == "h":
                self._flush()
        elif tag == "p":
            if self.kind in ("p", "quote"):
                self._flush()
        elif tag == "li":
            if self.kind == "li":
                self._flush()
        elif tag in ("ul", "ol"):
            self._flush()
            if self.lists:
                self.lists.pop()
        elif tag in ("strong", "b"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("em", "i"):
            self.ital = max(0, self.ital - 1)
        elif tag == "code" and not self.in_pre:
            self.code = max(0, self.code - 1)
        elif tag == "pre":
            self._flush()
            self.in_pre = False
        elif tag == "blockquote":
            self._flush()
            self.quote = max(0, self.quote - 1)
        elif tag == "a":
            if self.href and self.href.startswith(("http://", "https://")):
                self._add(f" ({self.href})", bold=False, ital=False, code=False)
            self.href = None
        elif tag in ("th", "td"):
            if self.row is not None and self.cell is not None:
                self.row.append((self.cell, self.cell_head, self.cell_align))
            self.cell = None
        elif tag == "tr":
            if self.table is not None and self.row is not None:
                self.table.append(self.row)
            self.row = None
        elif tag == "table":
            if self.table:
                self.blocks.append(("table", self.table))
            self.table = None

    def handle_data(self, data):
        if self.in_pre:
            self.runs.append((data, False, False, True))
            return
        text = re.sub(r"\s+", " ", data)
        if not text.strip() and (self.cell is None and not self.runs):
            return
        if text:
            self._add(text)


def _parse_markdown(text):
    p = _Blocks()
    p.feed(_md_to_html(text or ""))
    p.close()
    p._flush()
    return p.blocks


# ---------------------------------------------------------------------------
# Arculat a Word- és PDF-dokumentumokhoz (Contact-Line 2002 Bt. arculati kézikönyv, 2026. október):
# színek 0A2540 (Deep Corporate Blue), 00D4FF (Azure Light Blue, akcentus), F4F6F9 (Soft Background Gray);
# címsor Montserrat / Inter / Helvetica Bold, szöveg Open Sans / Roboto / Arial.
# Az első oldal a hivatalos levélpapír mintáját követi: balra a logó, jobbra a szlogen, alatta azúr vonal;
# a láblécben a cégadatok.
# Más arculathoz elég ezt a részt, a fonts/ mappa betűit és a logófájlt cserélni.
# ---------------------------------------------------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))
BRAND = {
    "primary": "0A2540",   # Deep Corporate Blue: címsorok, lábléc
    "dark": "0A2540",      # Deep Corporate Blue: szöveg, főcím
    "grey": "C9D3DF",      # a mélykék halvány árnyalata: táblázatvonalak, lábléc vonala
    "light": "F4F6F9",     # Soft Background Gray: táblázatfejléc háttere
    "accent": "00D4FF",    # Azure Light Blue: a fejléc alatti vonal (szövegszínnek túl világos)
    "logo": os.path.join(_HERE, "assets", "contact_line_logo.png"),
    "logo_width_cm": 6.2,
    "tagline": "Üzletviteli Tanácsadás és Szolgáltatás",
    "footer_text": "Contact-Line 2002 Bt. • Cégjegyzékszám: 01-06-782803 • Adószám: 21447013-2-41",
    "head_bold": True,     # a címsorbetű a félkövér változat (Helvetica Bold)
    # Word: a betűket a megnyitó gép adja, ezért itt csak a nevük szerepel (a kézikönyv tartalék betűi,
    # mert ezek minden gépen megvannak; a Helvetica helyett Windowson az Arial jelenik meg)
    "docx_body_font": "Arial",
    "docx_bold_font": "Arial",
    "docx_head_font": "Helvetica",
}
_HU_MONTHS = ("január", "február", "március", "április", "május", "június", "július", "augusztus",
              "szeptember", "október", "november", "december")


def _today_hu():
    """A mai dátum magyarul, budapesti idő szerint (a láblécbe)."""
    try:
        from zoneinfo import ZoneInfo
        d = datetime.datetime.now(ZoneInfo("Europe/Budapest")).date()
    except Exception:
        d = datetime.date.today()
    return f"{d.year}. {_HU_MONTHS[d.month - 1]} {d.day}."


def _logo(brand=None):
    """A logó (fájlút vagy bájtfolyam) és oldalaránya (magasság/szélesség); (None, 0), ha nincs vagy nem olvasható."""
    src = (brand or BRAND).get("logo")
    if not src:
        return None, 0
    try:
        if isinstance(src, (bytes, bytearray)):
            head = bytes(src[:24])
        else:
            if not os.path.exists(src):
                return None, 0
            with open(src, "rb") as fh:
                head = fh.read(24)
        if head[:8] != b"\x89PNG\r\n\x1a\n":
            return None, 0
        w, h = int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")
        if not (w and h):
            return None, 0
        return (io.BytesIO(bytes(src)) if isinstance(src, (bytes, bytearray)) else src), h / w
    except OSError:
        return None, 0


def make_brand(name="", color="333033", logo_png=None):
    """Egyszerű, semleges arculat (más cégekhez vagy üres sablonhoz): saját logó és kiemelőszín, semleges betűk.
    Logó nélkül üres fejléc."""
    color = (color or "333033").lstrip("#").upper()
    return {"name": name, "primary": color, "dark": "222222", "grey": "B5B5B5", "light": "F2F2F2",
            "logo": logo_png or None, "logo_width_cm": 3.4, "neutral": True, "accent": None, "tagline": "",
            "footer_text": "", "head_bold": False,
            "docx_body_font": "Calibri", "docx_bold_font": "Calibri", "docx_head_font": "Calibri"}


EMPTY_BRAND = make_brand("Üres")


def _clear_body(doc):
    """A fejléces papír törzsének kiürítése (a szakaszbeállítás – fejléc, lábléc, margó – marad)."""
    from docx.oxml.ns import qn
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _ensure_styles(doc):
    """A Sara-dokumentumhoz kellő stílusok pótlása a fejléces papírba, ha hiányoznak; a címsorok a sablon
    alapbetűjével (Normal) és félkövéren, a beépített kék szín nélkül."""
    import copy
    from docx import Document
    from docx.oxml.ns import qn
    have = {s.name for s in doc.styles}
    for s in doc.styles:   # a Word gyári (kék, témabetűs) címsorai helyett a papír alapbetűje, sötét színnel
        if s.name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
            rpr = s.element.find(qn("w:rPr"))
            if rpr is None:
                continue
            for x in rpr.findall(qn("w:color")):
                if x.get(qn("w:themeColor")):
                    rpr.remove(x)
            for x in rpr.findall(qn("w:rFonts")):
                if x.get(qn("w:asciiTheme")) and x.get(qn("w:asciiTheme")).startswith("major"):
                    rpr.remove(x)
    default = None
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Quote", "Title"):
        if name in have:
            continue
        default = default or Document()
        el = copy.deepcopy(default.styles[name].element)
        rpr = el.find(qn("w:rPr"))
        if rpr is not None:
            for tag in ("w:rFonts", "w:color"):
                for x in rpr.findall(qn(tag)):
                    rpr.remove(x)
        doc.styles.element.append(el)


def apply_letterhead(doc_bytes, letterhead_bytes):
    """Kész Word-dokumentum (pl. iratmintából kitöltött) átültetése fejléces papírra: a papír fejléce, lábléce
    (képekkel együtt) és oldalbeállítása kerül a dokumentum minden szakaszába; a törzs változatlan."""
    import copy
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(doc_bytes))
    lh = Document(io.BytesIO(letterhead_bytes))
    ls = lh.sections[0]
    rid_attrs = [qn("r:embed"), qn("r:id"), qn("r:link")]
    for sec in doc.sections:
        for attr in ("page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin",
                     "header_distance", "footer_distance"):
            value = getattr(ls, attr)
            if value is not None:
                setattr(sec, attr, value)
        sec.different_first_page_header_footer = ls.different_first_page_header_footer
        for name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header",
                     "even_page_footer"):
            src, dst = getattr(ls, name), getattr(sec, name)
            if src.is_linked_to_previous:
                dst.is_linked_to_previous = True
                continue
            dst.is_linked_to_previous = False
            el = dst._element
            for child in list(el):
                el.remove(child)
            for child in src._element:
                el.append(copy.deepcopy(child))
            for node in el.iter():
                for attr in rid_attrs:
                    rid = node.get(attr)
                    if not rid or rid not in src.part.rels:
                        continue
                    rel = src.part.rels[rid]
                    if rel.is_external:
                        node.set(attr, dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True))
                    elif rel.reltype.endswith("/image"):
                        new_rid, _ = dst.part.get_or_add_image(io.BytesIO(rel.target_part.blob))
                        node.set(attr, new_rid)
    doc.settings.odd_and_even_pages_header_footer = lh.settings.odd_and_even_pages_header_footer
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Markdown -> DOCX
# ---------------------------------------------------------------------------
def markdown_to_docx(md_text, title="", brand=None, base=None):
    """Markdown → Word. A base egy fejléces papír (.docx bájtjai): ekkor annak fejléce, lábléce, margói és betűi
    maradnak, a tartalom a törzsébe kerül; különben a Contact-Line 2002 (vagy a brand) arculata."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    B = {**BRAND, **(brand or {})}
    blocks = _parse_markdown(md_text)
    doc = Document(io.BytesIO(base)) if base else Document()
    if base:
        _clear_body(doc)
        _ensure_styles(doc)
    doc.core_properties.title = title or ""
    section = doc.sections[0]
    if not base:
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
        section.header_distance, section.footer_distance = Cm(1.2), Cm(1.0)
    text_width = section.page_width - section.left_margin - section.right_margin

    def rgb(key):
        return RGBColor.from_string(B[key])

    def set_font(style_or_run, name):
        """A betűtípus beállítása minden írásrendszerre; a sablon témabetűit felülírja."""
        rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)

    def border(par, side, color, size="4", space="4"):
        ppr = par._p.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        el = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:sz", size), ("w:space", space), ("w:color", color)):
            el.set(qn(k), v)
        pbdr.append(el)

    if not base:   # fejléces papírnál a sablon stílusai maradnak
        normal = doc.styles["Normal"]
        set_font(normal, B["docx_body_font"])
        normal.font.size, normal.font.color.rgb = Pt(10), rgb("dark")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "hu-HU")  # helyesírás-ellenőrzés magyarul
        normal.element.get_or_add_rPr().append(lang)

        for level, size, color, before, after in ((1, 22, "dark", 14, 8), (2, 15, "primary", 14, 6),
                                                  (3, 12.5, "dark", 10, 4), (4, 10.5, "primary", 8, 3)):
            st = doc.styles[f"Heading {level}"]
            set_font(st, B["docx_bold_font"] if level == 4 else B["docx_head_font"])
            st.font.size, st.font.color.rgb = Pt(size), rgb(color)
            st.font.bold, st.font.italic = bool(B.get("head_bold")) or level == 4, False
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
            st.paragraph_format.line_spacing = 1.1
            st.paragraph_format.keep_with_next = True
        quote_style = doc.styles["Quote"]
        set_font(quote_style, B["docx_body_font"])
        quote_style.font.italic, quote_style.font.color.rgb = True, rgb("primary")
        quote_style.paragraph_format.left_indent = Cm(0.6)

    def add_runs(par, runs, size=None, force_bold=False, plain_bold=False):
        for text, b, i, c in runs:
            for k, part in enumerate(text.split("\n")):
                if k:
                    par.add_run().add_break()
                if not part:
                    continue
                r = par.add_run(part)
                if (b or force_bold) and not c:
                    r.bold = True
                    if not plain_bold and not base:  # a Light változatnak nincs saját félkövérje, ezért az alapcsaládé kell
                        set_font(r._r.get_or_add_rPr(), B["docx_bold_font"])
                if i:
                    r.italic = True
                if c:
                    set_font(r._r.get_or_add_rPr(), "Consolas")
                    r.font.size = Pt(9)
                if size and not c:
                    r.font.size = Pt(size)

    align = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}

    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            level = min(blk[1], 4)
            add_runs(doc.add_heading(level=level), blk[2], plain_bold=True)
        elif kind == "p":
            add_runs(doc.add_paragraph(), blk[1])
        elif kind == "quote":
            par = doc.add_paragraph(style="Quote")
            border(par, "left", B["grey"], size="12", space="8")
            add_runs(par, blk[1])
        elif kind == "li":
            _, depth, ordered, n, runs = blk
            if ordered:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 * depth + 0.3)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                p.add_run(f"{n}. ")
            elif base:   # a sablonban nem biztos, hogy van felsorolás-számozás: kézi felsorolásjel
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 * depth + 0.3)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.space_after = Pt(3)
                p.add_run("•  ")
            else:
                p = doc.add_paragraph(style="List Bullet" if depth <= 1 else f"List Bullet {min(depth, 3)}")
                p.paragraph_format.space_after = Pt(3)
            add_runs(p, runs)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_runs(p, blk[1])
        elif kind == "hr":
            border(doc.add_paragraph(), "bottom", B["grey"], size="6", space="1")
        elif kind == "table":
            rows = blk[1]
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            # csak vízszintes vonalak, az arculat szürkéjével
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "bottom", "insideH"):
                el = OxmlElement(f"w:{side}")
                for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", B["grey"])):
                    el.set(qn(k), v)
                borders.append(el)
            table._tbl.tblPr.append(borders)
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(3)
                    par.paragraph_format.line_spacing = 1.1
                    if ci < len(row):
                        runs, head, al = row[ci]
                        add_runs(par, runs, size=9, force_bold=head)
                        if al in align:
                            par.alignment = align[al]
                        if head:
                            shd = OxmlElement("w:shd")
                            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", B["light"])):
                                shd.set(qn(k), v)
                            cell._tc.get_or_add_tcPr().append(shd)
            doc.add_paragraph()

    if not base:
        # Fejléc: a logó csak az első oldalon; lábléc minden oldalon: dátum balra, oldalszám jobbra
        logo, _ratio = _logo(B)
        section.different_first_page_header_footer = True
        if logo:
            hp = section.first_page_header.paragraphs[0]
            hp.paragraph_format.space_after = Pt(10)
            hp.add_run().add_picture(logo, width=Cm(B["logo_width_cm"]))
            if B.get("tagline"):   # jobbra a szlogen, a logó alsó vonalához igazítva
                header_tabs = doc.styles["Header"].element.get_or_add_pPr().find(qn("w:tabs"))
                if header_tabs is not None:   # a sablon középső tabulátora nélkül a szlogen a jobb szélre kerül
                    header_tabs.getparent().remove(header_tabs)
                hp.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                tr = hp.add_run("\t" + B["tagline"])
                tr.font.size, tr.font.color.rgb = Pt(8.5), rgb("primary")
            if B.get("accent"):
                border(hp, "bottom", B["accent"], size="8", space="6")
        today = _today_hu()
        footer_tabs = doc.styles["Footer"].element.get_or_add_pPr().find(qn("w:tabs"))
        if footer_tabs is not None:  # a sablon középre és jobbra igazító tabulátorai nélkül csak a sajátunk marad
            footer_tabs.getparent().remove(footer_tabs)

        def footer_run(par, text=None, field=None, instr=None):
            run = par.add_run(text) if text else par.add_run()
            run.font.size, run.font.color.rgb = Pt(8), rgb("primary")
            if field:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), field)
                run._r.append(el)
            if instr:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
                run._r.append(el)

        for footer in (section.footer, section.first_page_footer):
            fp = footer.paragraphs[0]
            fp.paragraph_format.space_after = Pt(0)
            fp.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            border(fp, "top", B["grey"], size="4", space="6")
            footer_run(fp, (B["footer_text"] + "\t" + today + "  |  ") if B.get("footer_text") else today + "\t")
            footer_run(fp, field="begin")
            footer_run(fp, instr="PAGE")
            footer_run(fp, field="separate")
            footer_run(fp, "1")
            footer_run(fp, field="end")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markdown -> PDF (reportlab; a betűk a fájlba ágyazódnak, így a PDF mindenhol egyformán néz ki)
# ---------------------------------------------------------------------------
_FONTS = None


def _register_fonts():
    """A betűk betöltése; a használandó betűnevek szótárát adja vissza.
    Sorrend a kézikönyv szerint: szöveg Open Sans, ha nincs, Liberation Sans (Arial-méretű); címsor Montserrat,
    ha nincs, Inter. Ha ezek mind hiányoznak a fonts/ mappából, DejaVu betűvel készül a PDF."""
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    dirs = [os.path.join(_HERE, "fonts"), "/usr/share/fonts/truetype/dejavu"]

    def find(name):
        for d in dirs:
            path = os.path.join(d, name)
            if os.path.exists(path):
                return path
        return None

    def register(pairs):
        if any(find(f) is None for _, f in pairs):
            return False
        for reg, file in pairs:
            pdfmetrics.registerFont(TTFont(reg, find(file)))
        return True

    if not register((("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf"),
                     ("DejaVu-Italic", "DejaVuSans-Oblique.ttf"), ("DejaVu-BoldItalic", "DejaVuSans-BoldOblique.ttf"),
                     ("DejaVuMono", "DejaVuSansMono.ttf"))):
        raise ToolError("Hiányzik a betűtípusfájl a PDF-hez: fonts/DejaVuSans*.ttf")
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                  italic="DejaVu-Italic", boldItalic="DejaVu-BoldItalic")
    fonts = {"body": "DejaVu", "bold": "DejaVu-Bold", "italic": "DejaVu-Italic", "head": "DejaVu-Bold",
             "fallback": "DejaVu", "mono": "DejaVuMono"}
    for fam, prefix in (("OpenSans", "OpenSans"), ("Liberation", "LiberationSans")):   # szövegbetű
        if register(((fam, f"{prefix}-Regular.ttf"), (f"{fam}-Bold", f"{prefix}-Bold.ttf"),
                     (f"{fam}-Italic", f"{prefix}-Italic.ttf"), (f"{fam}-BoldItalic", f"{prefix}-BoldItalic.ttf"))):
            pdfmetrics.registerFontFamily(fam, normal=fam, bold=f"{fam}-Bold", italic=f"{fam}-Italic",
                                          boldItalic=f"{fam}-BoldItalic")
            fonts.update(body=fam, bold=f"{fam}-Bold", italic=f"{fam}-Italic", head=f"{fam}-Bold")
            break
    for fam, file in (("Montserrat-Bold", "Montserrat-Bold.ttf"), ("Inter-Bold", "Inter-Bold.ttf")):   # címsorbetű
        if register(((fam, file),)):
            fonts["head"] = fam
            break
    fonts["glyphs"] = {name: set(pdfmetrics.getFont(name).face.charToGlyph)
                       for name in {fonts["body"], fonts["head"]}}
    _FONTS = fonts
    return fonts


def markdown_to_pdf(md_text, title="", brand=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    B = {**BRAND, **(brand or {})}
    F = dict(_register_fonts())
    if B.get("neutral"):   # semleges arculat: a címsor is a szövegbetű félkövérje
        F["head"] = F["bold"]
    blocks = _parse_markdown(md_text)
    dark, primary = colors.HexColor("#" + B["dark"]), colors.HexColor("#" + B["primary"])
    grey, light = colors.HexColor("#" + B["grey"]), colors.HexColor("#" + B["light"])

    body = ParagraphStyle("body", fontName=F["body"], fontSize=9.5, leading=14.5, spaceAfter=6, textColor=dark)
    heads = {}
    for level, size, lead, color, before, after in ((1, 21, 25, dark, 12, 8), (2, 15, 19, primary, 12, 6),
                                                    (3, 12.5, 16, dark, 9, 4), (4, 10, 14, primary, 7, 3)):
        heads[level] = ParagraphStyle(f"h{level}", parent=body, fontSize=size, leading=lead, textColor=color,
                                      fontName=F["bold"] if level == 4 else F["head"],
                                      spaceBefore=before, spaceAfter=after, keepWithNext=1)
    quote = ParagraphStyle("quote", parent=body, fontName=F["italic"], textColor=primary, leftIndent=14)
    cell_styles = {a: ParagraphStyle(f"cell_{a}", parent=body, fontSize=8.5, leading=11.5, spaceAfter=0, alignment=al)
                   for a, al in (("left", 0), ("center", 1), ("right", 2))}

    def with_fallback(text, font):
        """A betűtípusból hiányzó jeleket (pl. nyilak, pipák) a DejaVu betűvel írja ki."""
        have = F["glyphs"].get(font)
        if not have:
            return escape(text)
        out, buf, missing = [], [], False

        def flush():
            if buf:
                t = escape("".join(buf))
                out.append(f'<font name="{F["fallback"]}">{t}</font>' if missing else t)
                buf.clear()
        for ch in text:
            miss = ord(ch) not in have and not ch.isspace()
            if miss != missing:
                flush()
                missing = miss
            buf.append(ch)
        flush()
        return "".join(out)

    def markup(runs, force_bold=False, font=None):
        out = []
        for text, b, i, c in runs:
            if c:
                t = escape(text).replace("\n", "<br/>")
                t = f'<font name="{F["mono"]}" size="8.5">{t.replace("  ", "&nbsp;&nbsp;")}</font>'
            else:
                t = with_fallback(text, font or F["body"]).replace("\n", "<br/>")
            if b or force_bold:
                t = f"<b>{t}</b>"
            if i:
                t = f"<i>{t}</i>"
            out.append(t)
        return "".join(out)

    logo, ratio = _logo(B)
    logo_w = B["logo_width_cm"] * cm
    logo_h = logo_w * ratio
    left, right = 2 * cm + 6, A4[0] - 2 * cm - 6  # a szövegtükör szélei (a keret 6 pontos belső margójával)
    avail = right - left
    story = [Spacer(1, logo_h + (0.6 if B.get('accent') else 0.3) * cm)] if logo else []
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            level = min(blk[1], 4)
            story.append(Paragraph(markup(blk[2], font=F["bold"] if level == 4 else F["head"]), heads[level]))
        elif kind == "p":
            story.append(Paragraph(markup(blk[1]), body))
        elif kind == "quote":
            story.append(Paragraph(markup(blk[1]), quote))
        elif kind == "li":
            _, depth, ordered, n, runs = blk
            style = ParagraphStyle("li", parent=body, leftIndent=18 * depth + 4, bulletIndent=18 * (depth - 1) + 4,
                                   spaceAfter=3, bulletFontName=F["body"], bulletColor=primary)
            story.append(Paragraph(markup(runs), style, bulletText=f"{n}." if ordered else "•"))
        elif kind == "code":
            story.append(Paragraph(markup(blk[1]), ParagraphStyle("code", parent=body, fontName=F["mono"],
                                                                  fontSize=8.5, leading=11, leftIndent=8)))
        elif kind == "hr":
            story.append(HRFlowable(width="100%", thickness=0.5, color=grey, spaceBefore=4, spaceAfter=8))
        elif kind == "table":
            rows = blk[1]
            ncols = max(len(r) for r in rows)
            data, weights = [], [6] * ncols
            for row in rows:
                line = []
                for ci in range(ncols):
                    if ci < len(row):
                        runs, head, al = row[ci]
                        plain = "".join(r[0] for r in runs)
                        weights[ci] = max(weights[ci], min(len(plain), 40))
                        line.append(Paragraph(markup(runs, force_bold=head), cell_styles.get(al or "left")))
                    else:
                        line.append("")
                data.append(line)
            has_head = any(c[1] for c in rows[0])
            total = sum(weights)
            tbl = Table(data, colWidths=[avail * w / total for w in weights], repeatRows=1 if has_head else 0)
            style = [("LINEABOVE", (0, 0), (-1, 0), 0.5, grey), ("LINEBELOW", (0, 0), (-1, -1), 0.5, grey),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                     ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
            if has_head:
                style.append(("BACKGROUND", (0, 0), (-1, 0), light))
            tbl.setStyle(TableStyle(style))
            story.extend([tbl, Spacer(1, 8)])

    today = _today_hu()

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setStrokeColor(grey)
        canvas.setLineWidth(0.5)
        canvas.line(left, 1.55 * cm, right, 1.55 * cm)
        canvas.setFont(F["body"], 7.5)
        canvas.setFillColor(primary)
        if B.get("footer_text"):
            canvas.drawString(left, 1.1 * cm, B["footer_text"])
            canvas.drawRightString(right, 1.1 * cm, f"{today}  |  {doc_.page}")
        else:
            canvas.drawString(left, 1.1 * cm, today)
            canvas.drawRightString(right, 1.1 * cm, str(doc_.page))
        canvas.restoreState()

    def first_page(canvas, doc_):
        if logo:
            from reportlab.lib.utils import ImageReader
            img = ImageReader(logo) if not isinstance(logo, str) else logo
            if not isinstance(logo, str):
                logo.seek(0)
            top = A4[1] - 1.6 * cm
            canvas.drawImage(img, left, top - logo_h, width=logo_w, height=logo_h, mask="auto")
            if B.get("tagline"):   # jobbra a szlogen a logó alsó vonalán
                canvas.setFont(F["body"], 8)
                canvas.setFillColor(primary)
                canvas.drawRightString(right, top - logo_h + 1, B["tagline"])
            if B.get("accent"):   # azúr vonal a fejléc alatt, mint a levélpapíron
                canvas.setStrokeColor(colors.HexColor("#" + B["accent"]))
                canvas.setLineWidth(1)
                canvas.line(left, top - logo_h - 0.3 * cm, right, top - logo_h - 0.3 * cm)
        footer(canvas, doc_)

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2.1 * cm, title=title or "", author="Sara")
    pdf.build(story or [Paragraph(" ", body)], onFirstPage=first_page, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Office fájlok szöveggé alakítása (hogy Sara lássa a tartalmukat)
# ---------------------------------------------------------------------------
def _clip(text):
    if len(text) > MAX_TEXT_FOR_CHAT:
        return text[:MAX_TEXT_FOR_CHAT] + "\n[... a tartalom itt csonkolva ...]"
    return text


def xlsx_to_text(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=False)  # a képletek szövegként maradnak meg
    out = []
    for ws in wb.worksheets:
        out.append(f"## Munkalap: {ws.title}")
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r > 300:
                out.append("[... további sorok kihagyva ...]")
                break
            if any(v is not None for v in row):
                out.append(f"{r}: " + " | ".join("" if v is None else str(v) for v in row))
    return _clip("\n".join(out))


def pptx_to_text(data):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"## {n}. dia")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Jegyzet: " + slide.notes_slide.notes_text_frame.text.strip())
    return _clip("\n".join(out))


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _roman(n):
    out = ""
    for value, letters in ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
                           (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")):
        while n >= value:
            out, n = out + letters, n - value
    return out


def _letters(n):
    out = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(97 + rem) + out
    return out


def _fmt_number(n, fmt):
    if fmt == "lowerLetter":
        return _letters(n)
    if fmt == "upperLetter":
        return _letters(n).upper()
    if fmt == "lowerRoman":
        return _roman(n)
    if fmt == "upperRoman":
        return _roman(n).upper()
    if fmt == "decimalZero":
        return f"{n:02d}"
    if fmt == "bullet":
        return "•"
    if fmt == "none":
        return ""
    return str(n)


class _DocxNumbering:
    """A Word automatikus számozásának (1., 1.1., a), iii.) visszaállítása, hogy a pontszámok a szövegben
    is szerepeljenek. Enélkül egy szerződés „5.2. pontja” a kinyert szövegben nem található meg."""

    def __init__(self, zf):
        import xml.etree.ElementTree as ET
        self.levels = {}      # numId -> {ilvl: (start, numFmt, lvlText)}
        self.counters = {}    # numId -> {ilvl: aktuális szám}
        self.style_num = {}   # styleId -> (numId, ilvl)
        try:
            root = ET.fromstring(zf.read("word/numbering.xml"))
        except KeyError:
            root = None
        if root is not None:
            abstract = {}
            for a in root.findall(_W + "abstractNum"):
                lv = {}
                for l in a.findall(_W + "lvl"):
                    ilvl = int(l.get(_W + "ilvl", "0"))
                    start = l.find(_W + "start")
                    fmt = l.find(_W + "numFmt")
                    text = l.find(_W + "lvlText")
                    lv[ilvl] = (int(start.get(_W + "val")) if start is not None else 1,
                                fmt.get(_W + "val") if fmt is not None else "decimal",
                                text.get(_W + "val") if text is not None else f"%{ilvl + 1}.")
                abstract[a.get(_W + "abstractNumId")] = lv
            for n in root.findall(_W + "num"):
                ref = n.find(_W + "abstractNumId")
                lv = dict(abstract.get(ref.get(_W + "val") if ref is not None else None, {}))
                for ov in n.findall(_W + "lvlOverride"):
                    so = ov.find(_W + "startOverride")
                    ilvl = int(ov.get(_W + "ilvl", "0"))
                    if so is not None and ilvl in lv:
                        lv[ilvl] = (int(so.get(_W + "val")),) + lv[ilvl][1:]
                self.levels[n.get(_W + "numId")] = lv
        try:
            styles = ET.fromstring(zf.read("word/styles.xml"))
        except KeyError:
            styles = None
        if styles is not None:
            based = {}
            for s in styles.findall(_W + "style"):
                sid = s.get(_W + "styleId")
                b = s.find(_W + "basedOn")
                if b is not None:
                    based[sid] = b.get(_W + "val")
                num = s.find(f"{_W}pPr/{_W}numPr")
                if num is not None:
                    nid, lvl = num.find(_W + "numId"), num.find(_W + "ilvl")
                    if nid is not None:
                        self.style_num[sid] = (nid.get(_W + "val"), int(lvl.get(_W + "val")) if lvl is not None else 0)
            for sid in list(based):   # az alapstílus számozását örökli
                seen, cur = set(), sid
                while cur and cur not in self.style_num and cur not in seen:
                    seen.add(cur)
                    cur = based.get(cur)
                if cur in self.style_num and sid not in self.style_num:
                    self.style_num[sid] = self.style_num[cur]

    def label(self, p):
        ppr = p.find(_W + "pPr")
        num_id, ilvl = None, 0
        if ppr is not None:
            num = ppr.find(_W + "numPr")
            if num is not None:
                nid, lvl = num.find(_W + "numId"), num.find(_W + "ilvl")
                num_id = nid.get(_W + "val") if nid is not None else None
                ilvl = int(lvl.get(_W + "val")) if lvl is not None else 0
            if num_id is None:
                style = ppr.find(_W + "pStyle")
                if style is not None and style.get(_W + "val") in self.style_num:
                    num_id, ilvl = self.style_num[style.get(_W + "val")]
        if not num_id or num_id == "0" or num_id not in self.levels:
            return ""
        lv = self.levels[num_id]
        counters = self.counters.setdefault(num_id, {})
        start = lv.get(ilvl, (1, "decimal", ""))[0]
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:   # a magasabb szint lépésekor az alszintek újraindulnak
            del counters[deeper]
        text = lv.get(ilvl, (1, "decimal", f"%{ilvl + 1}."))[2]

        def repl(m):
            k = int(m.group(1)) - 1
            fmt = lv.get(k, (1, "decimal", ""))[1]
            return _fmt_number(counters.get(k, lv.get(k, (1,))[0]), fmt)
        out = re.sub(r"%(\d)", repl, text).strip()
        return out + " " if out else ""


class _DocxComments:
    """A Word-megjegyzések (comments.xml): ki, mikor, mit írt, melyik szövegrészhez, és melyik melyikre válasz.
    A szövegben [M1], [M2] … jelöli a helyüket, a megjegyzések a szöveg végére kerülnek."""
    W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
    W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"

    def __init__(self, zf):
        import xml.etree.ElementTree as ET
        self.items, self.order, self.anchor, self.open = {}, [], {}, set()
        try:
            root = ET.fromstring(zf.read("word/comments.xml"))
        except KeyError:
            return
        para_of = {}
        for c in root.findall(_W + "comment"):
            cid = c.get(_W + "id")
            paras = c.findall(_W + "p")
            self.items[cid] = {"author": c.get(_W + "author") or "ismeretlen", "date": (c.get(_W + "date") or "")[:10],
                               "text": " ".join(t for t in (_docx_para_text(p).strip() for p in paras) if t),
                               "parent": None, "done": False}
            if paras and paras[-1].get(self.W14 + "paraId"):
                para_of[paras[-1].get(self.W14 + "paraId")] = cid
        try:   # válaszok és lezárt megjegyzések (újabb Word-változatok)
            ext = ET.fromstring(zf.read("word/commentsExtended.xml"))
            for e in ext.iter(self.W15 + "commentEx"):
                cid = para_of.get(e.get(self.W15 + "paraId"))
                if cid:
                    self.items[cid]["parent"] = para_of.get(e.get(self.W15 + "paraIdParent"))
                    self.items[cid]["done"] = e.get(self.W15 + "done") == "1"
        except KeyError:
            pass

    def label(self, cid):
        if cid not in self.order:
            self.order.append(cid)
        return f"M{self.order.index(cid) + 1}"

    def summary(self):
        lines = []
        for cid in self.order + [c for c in self.items if c not in self.order]:
            item = self.items.get(cid)
            if not item:
                continue
            head = f"[{self.label(cid)}] {item['author']}" + (f", {item['date']}" if item["date"] else "")
            if item["parent"] in self.items:
                head += f" (válasz erre: {self.label(item['parent'])})"
            elif self.anchor.get(cid, "").strip():
                quoted = " ".join(self.anchor[cid].split())
                head += f" – a megjelölt rész: „{quoted[:300]}{'…' if len(quoted) > 300 else ''}”"
            if item["done"]:
                head += " [megoldottként jelölve]"
            lines.append(f"{head}: {item['text']}")
        return lines


def _docx_para_text(p, comments=None):
    parts = []
    for el in p.iter():
        if comments is not None and el.tag == _W + "commentRangeStart":
            comments.open.add(el.get(_W + "id"))
        elif comments is not None and el.tag == _W + "commentRangeEnd":
            comments.open.discard(el.get(_W + "id"))
        elif comments is not None and el.tag == _W + "commentReference":
            cid = el.get(_W + "id")
            if cid in comments.items:
                parts.append(f" [{comments.label(cid)}]")
        if el.tag == _W + "t":
            parts.append(el.text or "")
            if comments is not None:
                for cid in comments.open:
                    comments.anchor[cid] = comments.anchor.get(cid, "") + (el.text or "")
        elif el.tag == _W + "tab":
            parts.append("\t")
        elif el.tag in (_W + "br", _W + "cr"):
            parts.append("\n")
        # a törölt szöveg (w:delText) szándékosan kimarad: a dokumentum mostani állapota számít
    return "".join(parts)


def docx_to_text(data):
    """A Word-dokumentum teljes szövege a dokumentum sorrendjében: bekezdések automatikus számozással,
    táblázatok a helyükön, tartalomvezérlők és követett módosítások beszúrásai is; a végén a lábjegyzetek.
    (A python-docx bekezdéslistája ezek egy részét kihagyja, ezért közvetlenül az XML-t olvassuk.)"""
    import xml.etree.ElementTree as ET
    import zipfile
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        body = ET.fromstring(zf.read("word/document.xml")).find(_W + "body")
        numbering = _DocxNumbering(zf)
    except Exception:
        from docx import Document   # tartalék: a régi, egyszerűbb kinyerés
        doc = Document(io.BytesIO(data))
        out = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                out.append(" | ".join(c.text for c in row.cells))
        return _clip("\n".join(out))

    try:   # ha a megjegyzések hibásak, a szöveg attól még teljesen kiolvasható
        comments = _DocxComments(zf)
    except Exception:
        comments = _DocxComments.__new__(_DocxComments)
        comments.items, comments.order, comments.anchor, comments.open = {}, [], {}, set()
    out = []

    def walk(container):
        for el in container:
            if el.tag == _W + "p":
                text = _docx_para_text(el, comments)
                if text.strip():
                    out.append(numbering.label(el) + text)
            elif el.tag == _W + "tbl":
                for row in el.findall(_W + "tr"):
                    cells = []
                    for cell in row.findall(_W + "tc"):
                        cells.append(" ".join(t for t in (_docx_para_text(p, comments).strip()
                                                          for p in cell.iter(_W + "p")) if t))
                    if any(cells):
                        out.append(" | ".join(cells))
            elif el.tag == _W + "sdt":
                content = el.find(_W + "sdtContent")
                if content is not None:
                    walk(content)
            elif el.tag in (_W + "ins", _W + "customXml", _W + "smartTag"):
                walk(el)

    walk(body)
    try:
        notes = ET.fromstring(zf.read("word/footnotes.xml"))
        texts = [t for t in (_docx_para_text(n).strip() for n in notes.findall(_W + "footnote")) if t]
        if texts:
            out.append("\nLábjegyzetek:")
            out += [f"{i}. {t}" for i, t in enumerate(texts, 1)]
    except KeyError:
        pass
    notes = comments.summary()
    if notes:
        out.append("\nWord-megjegyzések (a szövegben [M1], [M2] … jelöli a helyüket):")
        out += notes
    return _clip("\n".join(out))


def office_to_text(mime, data):
    """Visszaad szöveget, ha a mime ismert irodai formátum, különben None."""
    if mime == XLSX_MIME:
        return xlsx_to_text(data)
    if mime == PPTX_MIME:
        return pptx_to_text(data)
    if mime == DOCX_MIME:
        return docx_to_text(data)
    return None


def pdf_to_text(data, max_chars=300_000):
    """Szöveges (vektoros) PDF tartalma. Szkennelt PDF-ből nem jön ki szöveg, ilyenkor üres a válasz."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out, total = [], 0
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            out.append(text)
            total += len(text)
            if total >= max_chars:
                break
    return "\n\n".join(out)[:max_chars]


def extract_text(name, mime, data):
    """A fájl szövege a kereséshez. None, ha a fájlból nem nyerhető ki szöveg (kép, sérült vagy védett fájl)."""
    name = str(name or "").lower()
    mime = str(mime or "")
    try:
        if mime.startswith("text/") or name.endswith((".txt", ".md", ".csv")):
            return data.decode("utf-8", errors="replace")
        if mime == PDF_MIME or name.endswith(".pdf"):
            return pdf_to_text(data)
        # a böngésző nem mindig küld pontos típust, ezért a kiterjesztés is dönthet
        by_ext = {".xlsx": XLSX_MIME, ".pptx": PPTX_MIME, ".docx": DOCX_MIME}
        mime = next((m for ext, m in by_ext.items() if name.endswith(ext)), mime)
        return office_to_text(mime, data)
    except Exception:
        return None
