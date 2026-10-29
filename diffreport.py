"""Verzió-összehasonlítás (Diff Report): két iratváltozat eltérései gépi pontossággal, és mellé a modell jogi
magyarázata arról, hogy az egyes módosítások kinek kedveznek és mit jelentenek.

Az eltéréseket a program határozza meg (difflib), nem a modell, így egyetlen módosítás sem maradhat ki és egy sem
lehet kitalált; a modell csak magyaráz. Készíthető változáskövetéses Word-irat (redline) is.
Streamlittől független, ezért tesztelhető.
"""
import difflib
import io
import json
import re
from datetime import datetime, timezone

MAX_CHANGES = 150

TOOL = {
    "name": "compare_versions",
    "description": (
        "Verzió-összehasonlítás (Diff Report): egy irat két változatának (pl. a saját tervezet és az ellenérdekű "
        "fél által visszaküldött változat) eltérései pontokra lebontva, szó szerint, és mindegyikhez jogi "
        "magyarázat: mit változtat, kinek kedvez, mennyire jelentős. Az eltéréseket a program gépi pontossággal "
        "állapítja meg. Akkor használd, ha a felhasználó két változat összevetését, a változások értékelését kéri."),
    "input_schema": {"type": "object", "properties": {
        "old_document": {"type": "string", "description": "A korábbi (kiinduló) változat neve."},
        "new_document": {"type": "string", "description": "Az újabb (módosított) változat neve."},
        "client_role": {"type": "string", "description": (
            "Kinek az oldalán áll a felhasználó (pl. „vevő”); ha nem tudod, kérdezd meg, vagy hagyd el, és semlegesen értékel.")},
        "context": {"type": "string", "description": "Nem kötelező: az ügy háttere a beszélgetésből."}},
        "required": ["old_document", "new_document"]},
}

REDLINE_TOOL = {
    "name": "create_redline_document",
    "description": (
        "Változáskövetéses Word-irat (redline) két iratváltozatból: az újabb változat szövege, amelyben a "
        "korábbihoz képesti törlések és beszúrások Word „Változások követése” jelöléssel látszanak, így a felhasználó "
        "egyenként elfogadhatja vagy elutasíthatja őket. Csak kérésre készítsd."),
    "input_schema": {"type": "object", "properties": {
        "old_document": {"type": "string", "description": "A korábbi változat neve."},
        "new_document": {"type": "string", "description": "Az újabb változat neve."},
        "title": {"type": "string", "description": "A fájl neve kiterjesztés nélkül (nem kötelező)."}},
        "required": ["old_document", "new_document"]},
}

SYSTEM = (
    "Tapasztalt magyar ügyvéd vagy (társasági és ingatlanjog). Megkapod egy irat két változata közötti, "
    "gépi úton megállapított eltéréseket sorszámmal. Mindegyikről döntsd el: mit jelent jogilag, kinek kedvez "
    "(a megbízónak, az ellenérdekű félnek, vagy semleges), és mennyire jelentős (LÉNYEGES / KISEBB / "
    "SZERKESZTÉSI – az utóbbi az elírás, számozás, központozás). Ne találj ki új eltérést, csak a megadottakról "
    "írj; az indoklás legyen konkrét és rövid (1–3 mondat). Ha egy változás kockázatos a megbízónak, írj "
    "javaslatot is (elfogadható / tárgyalandó / elutasítandó, és miért). Csak egy JSON-objektumot adj vissza:\n"
    '{"summary": "2–4 mondatos összkép: mi a módosítások iránya", "changes": [{"id": 1, "significance": '
    '"LÉNYEGES|KISEBB|SZERKESZTÉSI", "favours": "megbízó|ellenérdekű fél|semleges", "explanation": "", '
    '"recommendation": ""}]}')

SIG_ORDER = {"LÉNYEGES": 0, "KISEBB": 1, "SZERKESZTÉSI": 2}
REF = re.compile(r"^\s*((?:\d{1,3}\.)+(?:\d{1,3})?\.?/?|[IVX]{1,6}\.|\d{1,3}\)|[a-z]\))\s*")


def paragraphs(text):
    return [" ".join(p.split()) for p in str(text or "").splitlines() if p.strip()]


def _words(text):
    return re.findall(r"\S+\s*", text)


def inline_diff(old, new, context=8):
    """Szavankénti eltérés jelölve: [-törölt-] {+beszúrt+}; a hosszú változatlan részekből csak a változás
    körüli néhány szó marad (…)."""
    a, b = _words(old), _words(new)
    ops = difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
    out = []
    for k, (op, i1, i2, j1, j2) in enumerate(ops):
        if op == "equal":
            words = a[i1:i2]
            first, last = k == 0, k == len(ops) - 1
            if len(words) > 2 * context + 1 and not (first and last):
                if first:
                    words = ["… "] + words[-context:]
                elif last:
                    words = words[:context] + ["…"]
                else:
                    words = words[:context] + ["… "] + words[-context:]
            out.append("".join(words))
        else:
            if i2 > i1:
                out.append("[-" + "".join(a[i1:i2]).strip() + "-] ")
            if j2 > j1:
                out.append("{+" + "".join(b[j1:j2]).strip() + "+} ")
    return "".join(out).strip()


def _pair(olds, news):
    """Egy cserélt szakasz bekezdéseinek párosítása hasonlóság szerint: [(régi|None, új|None)]."""
    pairs, used = [], set()
    for o in olds:
        best, best_r = None, 0.45
        for k, n in enumerate(news):
            if k in used:
                continue
            r = difflib.SequenceMatcher(None, o, n, autojunk=False).ratio()
            if r > best_r:
                best, best_r = k, r
        if best is None:
            pairs.append((o, None))
        else:
            used.add(best)
            pairs.append((o, news[best]))
    # a beszúrt (párosítatlan) új bekezdések a helyükre
    out, ni = [], 0
    matched = {n for _, n in pairs if n is not None}
    for o, n in pairs:
        while n is not None and ni < len(news) and news[ni] != n:
            if news[ni] not in matched:
                out.append((None, news[ni]))
            ni += 1
        out.append((o, n))
        if n is not None:
            ni += 1
    out += [(None, n) for n in news[ni:] if n not in matched]
    return out


def compare(old_text, new_text):
    """A két változat összevetése. Visszaad: (sorozat, változások).
    sorozat: [(„equal”|„changed”|„deleted”|„inserted”, régi, új)] a redline-hoz;
    változások: [{"id", "kind", "ref", "old", "new", "diff"}]."""
    a, b = paragraphs(old_text), paragraphs(new_text)
    seq, changes, ref = [], [], ""
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "equal":
            for p in b[j1:j2]:
                seq.append(("equal", p, p))
                m = REF.match(p)
                ref = m.group(1) if m else ref
            continue
        for o, n in _pair(a[i1:i2], b[j1:j2]):
            m = REF.match(n or o or "")
            here = m.group(1) if m else ref
            if o is not None and n is not None:
                if " ".join(o.split()) == " ".join(n.split()):
                    seq.append(("equal", o, n))
                    continue
                seq.append(("changed", o, n))
                changes.append({"kind": "módosítás", "ref": here, "old": o, "new": n, "diff": inline_diff(o, n)})
            elif o is not None:
                seq.append(("deleted", o, None))
                changes.append({"kind": "törlés", "ref": here, "old": o, "new": "", "diff": "[-" + o + "-]"})
            else:
                seq.append(("inserted", None, n))
                changes.append({"kind": "új szöveg", "ref": here, "old": "", "new": n, "diff": "{+" + n + "+}"})
            if n is not None and m:
                ref = m.group(1)
    for k, c in enumerate(changes, 1):
        c["id"] = k
    return seq, changes


def _text_of(response):
    return "".join(getattr(b, "text", "") for b in getattr(response, "content", [])
                   if getattr(b, "type", "") == "text").strip()


def explain(client, model, changes, old_name, new_name, client_role="", context="", log=None):
    """A modell magyarázata a változásokhoz: {id: {...}}, és az összkép."""
    if not changes:
        return "", {}
    listing = "\n".join(f"#{c['id']} ({c['kind']}{', ' + c['ref'] if c['ref'] else ''}): {c['diff'][:1500]}"
                        for c in changes[:MAX_CHANGES])
    user = (f"Korábbi változat: {old_name}\nÚjabb változat: {new_name}\n"
            f"A megbízó: {client_role or 'nincs megadva – értékelj semlegesen'}\n"
            + (f"Az ügy háttere: {context}\n" if context else "")
            + "\nAz eltérések ([-törölt-], {+beszúrt+}):\n" + listing)
    kwargs = dict(model=model, max_tokens=8000, system=SYSTEM, messages=[{"role": "user", "content": user}])
    try:
        response = client.messages.create(thinking={"type": "disabled"}, **kwargs)
    except Exception as e:
        if getattr(e, "status_code", None) != 400 or "thinking" not in str(e).lower():
            raise
        response = client.messages.create(**kwargs)
    if log:
        log(model, getattr(response, "usage", None))
    raw = _text_of(response)
    m = re.search(r"\{.*\}", raw, re.S)
    try:
        data = json.loads(m.group(0)) if m else {}
    except ValueError:
        data = {}
    by_id = {}
    for x in data.get("changes") or []:
        try:
            by_id[int(x.get("id"))] = x
        except (TypeError, ValueError):
            continue
    return str(data.get("summary") or "").strip(), by_id


def report(client, model, old_name, old_text, new_name, new_text, client_role="", context="", log=None):
    seq, changes = compare(old_text, new_text)
    if not paragraphs(old_text) or not paragraphs(new_text):
        raise ValueError("Az egyik változat szövege nem olvasható ki, ezért nem hasonlítható össze.")
    head = [f"VERZIÓ-ÖSSZEHASONLÍTÁS – {old_name} → {new_name}"]
    if client_role:
        head.append(f"Szemszög: {client_role}")
    if not changes:
        return "\n".join(head + ["", "A két változat szövege megegyezik (a formázást nem vizsgáltam)."])
    summary, notes = explain(client, model, changes, old_name, new_name, client_role, context, log)
    for c in changes:
        n = notes.get(c["id"], {})
        sig = str(n.get("significance", "")).upper()
        c["significance"] = sig if sig in SIG_ORDER else "KISEBB"
        c["favours"] = str(n.get("favours") or "").strip()
        c["explanation"] = str(n.get("explanation") or "").strip()
        c["recommendation"] = str(n.get("recommendation") or "").strip()
    counts = {s: sum(1 for c in changes if c["significance"] == s) for s in SIG_ORDER}
    lines = head + [f"Eltérések: {len(changes)} ({counts['LÉNYEGES']} lényeges, {counts['KISEBB']} kisebb, "
                    f"{counts['SZERKESZTÉSI']} szerkesztési)", ""]
    if summary:
        lines += ["Összkép: " + summary, ""]
    for c in sorted(changes, key=lambda c: (SIG_ORDER[c["significance"]], c["id"])):
        lines.append(f"#{c['id']} [{c['significance']}] {c['kind']}" + (f" – {c['ref']}" if c["ref"] else "")
                     + (f" – kinek kedvez: {c['favours']}" if c["favours"] else ""))
        lines.append(f"   Eltérés: {c['diff'][:1200]}")
        if c["explanation"]:
            lines.append(f"   Jelentése: {c['explanation']}")
        if c["recommendation"]:
            lines.append(f"   Javaslat: {c['recommendation']}")
    if len(changes) > MAX_CHANGES:
        lines += ["", f"Megjegyzés: {len(changes)} eltérés van; az első {MAX_CHANGES}-hoz készült magyarázat."]
    lines += ["", "Jelölés: [-törölt szöveg-] {+beszúrt szöveg+}. Az eltéréseket a program gépi pontossággal "
                  "állapította meg; a formázási (betű, igazítás) eltéréseket nem vizsgálja.",
              "A jelentést a lényeges tételekkel kezdve mutasd meg a felhasználónak. Változáskövetéses Word-iratot "
              "(redline) kérésre a create_redline_document eszközzel készíthetsz."]
    return "\n".join(lines)


# --- REDLINE (változáskövetéses Word) ---------------------------------------------------------------------------
def redline_docx(old_text, new_text, title="", author="Sara"):
    """Az újabb változat Word-iratként, a korábbihoz képesti eltérések követett változásként (w:ins / w:del)."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    seq, changes = compare(old_text, new_text)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Century Gothic"
    style.font.size = Pt(10)
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    rev = [0]

    def run(text, deleted=False):
        r = OxmlElement("w:r")
        t = OxmlElement("w:delText" if deleted else "w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        return r

    def mark(tag, text):
        rev[0] += 1
        el = OxmlElement(tag)
        el.set(qn("w:id"), str(rev[0]))
        el.set(qn("w:author"), author)
        el.set(qn("w:date"), date)
        el.append(run(text, deleted=(tag == "w:del")))
        return el

    if title:
        h = doc.add_paragraph()
        h.add_run(title).bold = True
    for kind, old, new in seq:
        p = doc.add_paragraph()._p
        if kind == "equal":
            p.append(run(new))
        elif kind == "inserted":
            p.append(mark("w:ins", new))
        elif kind == "deleted":
            p.append(mark("w:del", old))
        else:
            a, b = _words(old), _words(new)
            for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
                if op == "equal":
                    p.append(run("".join(b[j1:j2])))
                    continue
                if i2 > i1:
                    p.append(mark("w:del", "".join(a[i1:i2])))
                if j2 > j1:
                    p.append(mark("w:ins", "".join(b[j1:j2])))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), len(changes)
